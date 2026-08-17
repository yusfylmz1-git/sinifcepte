import os
import sys
import re
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

desktop = r"C:\Users\Okul\Desktop"
target_dir = None
for d in os.listdir(desktop):
    if "Kazan" in d or "kazan" in d.lower():
        target_dir = os.path.join(desktop, d)
        break

patterns = {}

for root, dirs, files in os.walk(target_dir):
    for f in files:
        if f.endswith(".xlsx") and not f.startswith("~$"):
            path = os.path.join(root, f)
            try:
                wb = openpyxl.load_workbook(path, data_only=True)
                for s in wb.sheetnames:
                    if s.lower() == "sayfa1" or wb[s].max_row <= 1:
                        continue
                    sheet = wb[s]
                    
                    # Find header row
                    header_row_idx = -1
                    headers = []
                    for r in range(1, min(10, sheet.max_row + 1)):
                        vals = [str(c.value).strip().replace('\n', ' ') for c in sheet[r] if c.value is not None]
                        u_line = " ".join(vals).upper()
                        if "HAFTA" in u_line or "SÜRE" in u_line or "DURATION" in u_line or "WEEK" in u_line or "ZEIT" in u_line:
                            header_row_idx = r
                            headers = [str(sheet.cell(r, c).value or '').replace('\n', ' ').strip() for c in range(1, sheet.max_column + 1)]
                            break
                    
                    # First data row
                    sample_data = []
                    if header_row_idx != -1:
                        for r in range(header_row_idx + 1, min(header_row_idx + 5, sheet.max_row + 1)):
                            row_vals = [str(sheet.cell(r, c).value or '').replace('\n', ' ').strip() for c in range(1, sheet.max_column + 1)]
                            if any(row_vals):
                                sample_data.append(row_vals)
                    
                    print(f"\n=======================================================")
                    print(f"FILE: {f[:35]} | SHEET: {s}")
                    print(f"Header Row (R{header_row_idx}):")
                    for c_idx, h in enumerate(headers, 1):
                        if h:
                            print(f"   Col {c_idx:02d}: [{h}]")
                    if sample_data:
                        print("Sample Row 1 Data:")
                        for c_idx, val in enumerate(sample_data[0], 1):
                            if val:
                                print(f"   Col {c_idx:02d}: [{val[:70]}]")
            except Exception as e:
                print(f"Error on {f}: {e}")
