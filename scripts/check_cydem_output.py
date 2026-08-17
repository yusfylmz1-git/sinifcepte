import openpyxl, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r"C:\Users\Okul\Desktop\TÜM_DERSLER_MERKEZ_STANDART.xlsx", data_only=True)
for s in wb.sheetnames:
    if "7.Snf İngilizce (ÇYDEM)" in s or "7.Snf İngilizce" in s:
        sheet = wb[s]
        print(f"\n==================== SHEET: {s} ====================")
        for r in range(1, 18):
            row_vals = [str(sheet.cell(r, c).value or '') for c in range(1, 8)]
            print(f"R{r:02d}: " + " | ".join(row_vals))
        break
