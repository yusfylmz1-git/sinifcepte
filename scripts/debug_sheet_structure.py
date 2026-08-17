import os
import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\Okul\Desktop\Kazanımlar"

for f in os.listdir(folder):
    if f.endswith(".xlsx") and not f.startswith("~$"):
        path = os.path.join(folder, f)
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            print(f"\n==============================\nFILE: {f}")
            for s in wb.sheetnames:
                sheet = wb[s]
                print(f"\n--- Sheet: {s} (rows={sheet.max_row}, cols={sheet.max_column}) ---")
                for r in range(1, min(10, sheet.max_row + 1)):
                    row_vals = []
                    for c_idx, cell in enumerate(sheet[r], start=1):
                        if cell.value is not None:
                            val_str = str(cell.value).replace('\n', ' ')
                            row_vals.append(f"C{c_idx}[{val_str[:35]}]")
                    if row_vals:
                        print(f"Row {r:02d}: " + " | ".join(row_vals[:5]))
        except Exception as e:
            print(f"Error on {f}: {e}")

