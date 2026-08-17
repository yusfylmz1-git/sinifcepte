import sys
import os
import re
import json
import openpyxl

# Set utf-8 encoding for stdout
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

FOLDER = r"C:\Users\Okul\Desktop\Kazanımlar"
OUTPUT_JSON = r"c:\Users\Okul\Desktop\Projelerim\sinifcepte\admin_portal\data\official_maarif_kazanimlar.json"
PRESETS_JS = r"c:\Users\Okul\Desktop\Projelerim\sinifcepte\admin_portal\js\curriculum_presets.js"

os.makedirs(os.path.dirname(OUTPUT_JSON), exist_ok=True)

def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    text = re.sub(r"\s+", " ", text)
    return text

def parse_grade(text):
    match = re.search(r"(\d+)\s*\.?\s*s[ıi]n[ıi]f", text, re.I)
    if match:
        return int(match.group(1))
    match2 = re.search(r"\b([1-9]|1[0-2])\b", text)
    if match2:
        return int(match2.group(1))
    return None

def detect_subject(filename, sheetname):
    combined = (filename + " " + sheetname).lower()
    if "matematik" in combined:
        return "MAT", "Matematik"
    if "fen" in combined:
        return "FEN", "Fen Bilimleri"
    if "türkçe" in combined or "turkce" in combined:
        return "TURKCE", "Türkçe"
    if "sosyal" in combined:
        return "SOSYAL", "Sosyal Bilgiler"
    if "inkılap" in combined or "inkilap" in combined or "inkılâp" in combined:
        return "INKILAP", "T.C. İnkılap Tarihi ve Atatürkçülük"
    if "hayat" in combined:
        return "HAYAT", "Hayat Bilgisi"
    if "ingilizce" in combined or "english" in combined:
        return "INGILIZCE", "İngilizce"
    if "bilişim" in combined or "bilisim" in combined:
        return "BILISIM", "Bilişim Teknolojileri ve Yazılım"
    if "din" in combined:
        return "DIN", "Din Kültürü ve Ahlak Bilgisi"
    if "beden" in combined:
        return "BEDEN", "Beden Eğitimi ve Spor"
    if "müzik" in combined or "muzik" in combined:
        return "MUZIK", "Müzik"
    if "görsel" in combined or "gorsel" in combined:
        return "GORSEL", "Görsel Sanatlar"
    if "almanca" in combined:
        return "ALMANCA", "Almanca"
    return "GENEL", "Genel Ders"

def parse_sheet(sheet, filename, sheetname):
    sub_code, sub_name = detect_subject(filename, sheetname)
    grade = parse_grade(sheetname) or parse_grade(filename) or 5

    outcomes = []
    current_week = 0
    current_unit = "1. Ünite / Tema"
    current_topic = "Genel Konular"

    for r in range(1, sheet.max_row + 1):
        row = sheet[r]
        row_vals = [clean_text(cell.value) for cell in row]
        if not any(row_vals):
            continue

        full_row_str = " ".join(row_vals)

        # Hafta numarasını satırdaki hücrelerden yakala
        found_week = None
        week_col_idx = None
        for idx, val in enumerate(row_vals[:4]):
            m = re.search(r"(\d{1,2})\s*\.?\s*(?:Hafta|Woche|Week)", val, re.I)
            if m:
                found_week = int(m.group(1))
                week_col_idx = idx
                break

        if found_week is not None:
            current_week = found_week

        if current_week == 0 or current_week > 38:
            continue

        potential_texts = [v for idx, v in enumerate(row_vals) if v and idx != week_col_idx and not re.match(r"^\d+$", v)]
        
        unit_cand = ""
        topic_cand = ""
        outcome_cand = ""

        for text in potential_texts:
            if "ÜNİTE" in text.upper() or "TEMA" in text.upper() or "THEME" in text.upper() or "ÖĞRENME ALANI" in text.upper():
                unit_cand = text
            elif re.search(r"([A-ZÇĞİÖŞÜa-zçğıöşü]{1,4}\.\d+\.\d+|\b[A-Z]\d\.\d\.\w+\b|Students will be able to)", text):
                outcome_cand = text
            elif len(text) > 10 and not outcome_cand:
                topic_cand = text

        if unit_cand: current_unit = unit_cand
        if topic_cand: current_topic = topic_cand
        if not outcome_cand and potential_texts:
            outcome_cand = max(potential_texts, key=len)

        code_match = re.search(r"([A-ZÇĞİÖŞÜa-zçğıöşü]{1,4}\.\d+\.\d+(?:\.\d+)?|\b[A-Z]\d\.\d\.\w+\b)", outcome_cand)
        outcome_code = code_match.group(1) if code_match else f"{sub_code}.{grade}.{current_week}"

        is_holiday = current_week in (9, 28)
        holiday_note = f"{1 if current_week == 9 else 2}. Dönem Ara Tatili" if is_holiday else None

        outcomes.append({
            "id": f"out_{grade}_{sub_code}_w{current_week}",
            "gradeLevel": grade,
            "subjectCode": sub_code,
            "subjectName": sub_name,
            "weekNumber": current_week,
            "unitTitle": current_unit,
            "topicTitle": current_topic,
            "outcomeCode": outcome_code,
            "outcomeDescription": outcome_cand or f"{sub_name} {current_week}. hafta MEB müfredat kazanımı ve etkinlikleri.",
            "isHolidayWeek": is_holiday,
            "holidayNote": holiday_note
        })

    dedup = {}
    for o in outcomes:
        w = o["weekNumber"]
        if w not in dedup:
            dedup[w] = o
        else:
            if len(o["outcomeDescription"]) > len(dedup[w]["outcomeDescription"]):
                dedup[w]["outcomeDescription"] = o["outcomeDescription"]

    res = []
    if dedup:
        for w in range(1, 37):
            if w in dedup:
                res.append(dedup[w])
            else:
                prev = res[-1] if res else None
                u_title = prev["unitTitle"] if prev else "1. Ünite"
                t_title = prev["topicTitle"] if prev else "Genel Uygulamalar"
                is_hol = w in (9, 28)
                res.append({
                    "id": f"out_{grade}_{sub_code}_w{w}",
                    "gradeLevel": grade,
                    "subjectCode": sub_code,
                    "subjectName": sub_name,
                    "weekNumber": w,
                    "unitTitle": f"{1 if w == 9 else 2}. Dönem Ara Tatili" if is_hol else u_title,
                    "topicTitle": "Ara Tatil Etkinlikleri" if is_hol else t_title,
                    "outcomeCode": "TATIL" if is_hol else f"{sub_code}.{grade}.{w}",
                    "outcomeDescription": f"{1 if w == 9 else 2}. Dönem Ara Tatil Haftası" if is_hol else f"{sub_name} {w}. hafta MEB müfredat kazanımı ve etkinlikleri.",
                    "isHolidayWeek": is_hol,
                    "holidayNote": f"{1 if w == 9 else 2}. Dönem Ara Tatili" if is_hol else None
                })

    return res

def main():
    all_outcomes = []
    subject_map = {}

    print("Parsing all Excel files in:", FOLDER)

    for root, dirs, files in os.walk(FOLDER):
        for f in files:
            if f.endswith('.xlsx') and not f.startswith('~$'):
                path = os.path.join(root, f)
                try:
                    wb = openpyxl.load_workbook(path, data_only=True)
                    for s in wb.sheetnames:
                        sheet_outcomes = parse_sheet(wb[s], f, s)
                        if sheet_outcomes and len(sheet_outcomes) >= 30:
                            key = f"{sheet_outcomes[0]['gradeLevel']}_{sheet_outcomes[0]['subjectCode']}"
                            if key not in subject_map:
                                subject_map[key] = sheet_outcomes
                                all_outcomes.extend(sheet_outcomes)
                                print(f"  OK: {sheet_outcomes[0]['gradeLevel']}. Sinif {sheet_outcomes[0]['subjectName']} ({len(sheet_outcomes)} hafta)")
                except Exception as e:
                    print("  Error processing:", f, str(e))

    with open(OUTPUT_JSON, "w", encoding="utf-8") as out:
        json.dump(all_outcomes, out, ensure_ascii=False, indent=2)

    print(f"\nTOPLAM {len(all_outcomes)} HAFTALIK RESMI KAZANIM CIKARILDI!")
    print(f"Toplam Brans / Kademe Sayisi: {len(subject_map)}")

    generate_presets_js(subject_map, all_outcomes)

def generate_presets_js(subject_map, all_outcomes):
    presets_json_str = json.dumps(all_outcomes, ensure_ascii=False)
    
    js_content = f"""/**
 * SınıfCepte Web Admin Paneli - Resmî MEB / Maarif Modeli Müfredat Kazanım Kütüphanesi
 * Masaüstündeki 2025-2026 Resmî Çerçeve Yıllık Planlarından Otomatik Üretilmiştir.
 */
class CurriculumPresets {{
  // Tüm MEB Çerçeve Yıllık Planları Veritabanı ({len(all_outcomes)} Hafta)
  static OFFICIAL_DATABASE = {presets_json_str};

  static PRESETS = {{
    'full_maarif_library': {{
      title: '🌟 Tüm Resmî MEB / Maarif Modeli Kütüphanesini Yükle (Full Kütüphane)',
      description: 'Masaüstünüzdeki tüm derslerin ({len(all_outcomes)} haftalık) 1-8. sınıf müfredatını tek tıkla yükler.',
      isFull: true,
    }},
    'ortaokul_matematik': {{
      title: '📐 Ortaokul Matematik Paketi (5, 6, 7, 8. Sınıf)',
      description: '5, 6, 7, 8. Sınıf 36 haftalık resmî MEB Matematik çerçeve yıllık planları.',
      filter: (o) => o.subjectCode === 'MAT' && [5,6,7,8].includes(o.gradeLevel),
    }},
    'ilkokul_matematik': {{
      title: '🔢 İlkokul Matematik Paketi (1, 2, 3, 4. Sınıf)',
      description: '1, 2, 3, 4. Sınıf 36 haftalık resmî MEB İlkokul Matematik planları.',
      filter: (o) => o.subjectCode === 'MAT' && [1,2,3,4].includes(o.gradeLevel),
    }},
    'ortaokul_fen': {{
      title: '🔬 Ortaokul Fen Bilimleri Paketi (5, 6, 7, 8. Sınıf)',
      description: '5, 6, 7, 8. Sınıf 36 haftalık resmî MEB Fen Bilimleri planları.',
      filter: (o) => o.subjectCode === 'FEN' && [5,6,7,8].includes(o.gradeLevel),
    }},
    'turkce_full': {{
      title: '📚 Türkçe Paketi (1-8. Sınıflar Tüm Kademeler)',
      description: '1, 2, 3, 4, 5, 6, 7, 8. Sınıf 36 haftalık resmî MEB Türkçe yıllık planları.',
      filter: (o) => o.subjectCode === 'TURKCE',
    }},
    'sosyal_inkilap': {{
      title: '🌍 Sosyal Bilgiler & T.C. İnkılap Tarihi (4-8. Sınıf)',
      description: '4, 5, 6, 7. Sınıf Sosyal Bilgiler ve 8. Sınıf İnkılap Tarihi resmî planları.',
      filter: (o) => ['SOSYAL', 'INKILAP'].includes(o.subjectCode),
    }},
    'ingilizce_full': {{
      title: '🇬🇧 İngilizce Paketi (2-8. Sınıflar Tüm Kademeler)',
      description: '2, 3, 4, 5, 6, 7, 8. Sınıf MEB ve İYDEM İngilizce çerçeve yıllık planları.',
      filter: (o) => o.subjectCode === 'INGILIZCE',
    }},
    'hayat_bilgisi': {{
      title: '🌱 Hayat Bilgisi Paketi (1, 2, 3. Sınıf)',
      description: '1, 2, 3. Sınıf resmî Hayat Bilgisi çerçeve yıllık planları.',
      filter: (o) => o.subjectCode === 'HAYAT',
    }},
    'bilisim_teknolojileri': {{
      title: '💻 Bilişim Teknolojileri ve Yazılım (5, 6. Sınıf)',
      description: '5 ve 6. Sınıf TYMM Bilişim Teknolojileri ve Yazılım planları.',
      filter: (o) => o.subjectCode === 'BILISIM',
    }},
    'beden_muzik_gorsel': {{
      title: '🎨 Görsel Sanatlar, Müzik, Beden Eğitimi & Spor',
      description: '1-8. Sınıf Görsel Sanatlar, Müzik, Beden Eğitimi ve Oyun/Spor planları.',
      filter: (o) => ['BEDEN', 'MUZIK', 'GORSEL'].includes(o.subjectCode),
    }}
  }};

  /**
   * Seçilen paketi veya full kütüphaneyi sisteme yükler
   */
  static applyPreset(presetKey, outcomesManager) {{
    const preset = this.PRESETS[presetKey];
    if (!preset) return 0;

    let toAdd = [];
    if (preset.isFull) {{
      toAdd = this.OFFICIAL_DATABASE;
    }} else if (preset.filter) {{
      toAdd = this.OFFICIAL_DATABASE.filter(preset.filter);
    }}

    if (toAdd.length === 0) return 0;

    // Yüklenecek branş ve sınıfları mevcut listeden temizle
    const keysToReplace = new Set(toAdd.map(o => `${{o.gradeLevel}}_${{o.subjectCode}}`));
    outcomesManager.outcomes = outcomesManager.outcomes.filter(
      (o) => !keysToReplace.has(`${{o.gradeLevel}}_${{o.subjectCode}}`)
    );

    // Yeni kazanımları ekle
    outcomesManager.outcomes.push(...toAdd);
    outcomesManager.save();

    return toAdd.length;
  }}
}}

window.CurriculumPresets = CurriculumPresets;
"""

    with open(PRESETS_JS, "w", encoding="utf-8") as js_file:
        js_file.write(js_content)
    print(f"Updated {PRESETS_JS} with official library presets!")

if __name__ == "__main__":
    main()
